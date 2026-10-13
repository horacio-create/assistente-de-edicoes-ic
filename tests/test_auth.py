import io
from PIL import Image
import json
import os
import sys
import tempfile
import unittest
import uuid
from datetime import datetime, timedelta
from unittest import mock
from openpyxl import load_workbook
from pathlib import Path

SANDBOX = tempfile.TemporaryDirectory(prefix='indoor-auth-')
os.environ['INDOOR_DATA'] = str(Path(SANDBOX.name) / 'data')
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import server  # noqa: E402
from modules import auth, clockify, relatorios  # noqa: E402
from modules import ofertas  # noqa: E402
from storage import init, connect, now  # noqa: E402

init()
auth.init()
ofertas.garantir()


def call(method, path, body=None, cookie=None):
    data = json.dumps(body).encode() if body is not None else b''
    path, _, qs = path.partition('?')
    env = {'REQUEST_METHOD': method, 'PATH_INFO': path, 'QUERY_STRING': qs, 'HTTP_HOST': 'localhost:8080',
           'CONTENT_LENGTH': str(len(data)), 'wsgi.input': io.BytesIO(data), 'HTTP_X_INDOOR': '1',
           **({'HTTP_COOKIE': cookie} if cookie else {})}
    out = {}
    raw = b''.join(server.app(env, lambda s, h: out.update(status=s, headers=dict(h))))
    return int(out['status'][:3]), (json.loads(raw) if raw and out['headers'].get('Content-Type', '').startswith('application/json') else raw), out['headers']


def session(headers): return headers['Set-Cookie'].split(';')[0]


def signup(email, name='Pessoa', role=None, superadmin=False):
    token = auth.create_invite(email, role, superadmin)
    status, user, headers = call('POST', '/api/auth/cadastro', {'token': token, 'name': name, 'password': 'senha-forte'})
    assert status == 200, user
    return user, session(headers)


class AuthTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        server.AUTH = True
        with connect() as db:  # trabalho feito antes do login existir: vai para o primeiro superadmin
            db.execute("INSERT INTO jobs(id,title,updated,meta) VALUES('legado','Antigo',?,'{}')", (now(),))
        cls.admin, cls.admin_cookie = signup('chefe@x.com', 'Chefe', superadmin=True)
        status, role, _ = call('POST', '/api/admin/cargo-salvar', {'name': 'Ofertas', 'modules': ['offers']}, cls.admin_cookie)
        assert status == 200
        cls.role = role['id']

    @classmethod
    def tearDownClass(cls): server.AUTH = False

    def test_without_session_api_is_401_and_page_redirects(self):
        self.assertEqual(call('GET', '/api/jobs')[0], 401)
        status, _, headers = call('GET', '/')
        self.assertEqual((status, headers['Location']), (302, '/login.html'))
        self.assertEqual(call('GET', '/login.html')[0], 200)
        self.assertEqual(call('GET', '/api/info')[1], {'version': server.VERSION, 'auth': True, 'user': None})

    def test_legacy_data_goes_to_first_superadmin(self):
        status, jobs, _ = call('GET', '/api/jobs', cookie=self.admin_cookie)
        self.assertEqual(status, 200)
        self.assertIn('legado', [j['id'] for j in jobs])
        with connect() as db: self.assertIsNone(db.execute('SELECT 1 FROM jobs WHERE owner IS NULL').fetchone())

    def test_invite_is_single_use_and_login_works(self):
        token = auth.create_invite('uso@x.com', self.role)
        self.assertEqual(call('GET', '/api/auth/convite?token=' + token)[1], {'email': 'uso@x.com', 'role': 'Ofertas', 'reset': False})
        self.assertEqual(call('POST', '/api/auth/cadastro', {'token': token, 'name': 'Uso', 'password': 'senha-forte'})[0], 200)
        self.assertEqual(call('POST', '/api/auth/cadastro', {'token': token, 'name': 'Uso', 'password': 'senha-forte'})[0], 400)
        status, user, headers = call('POST', '/api/auth/login', {'email': 'USO@x.com ', 'password': 'senha-forte'})
        self.assertEqual((status, user['email']), (200, 'uso@x.com'))
        self.assertEqual(call('GET', '/api/auth/me', cookie=session(headers))[1]['name'], 'Uso')
        self.assertEqual(call('POST', '/api/auth/logout', {}, session(headers))[0], 200)
        self.assertEqual(call('GET', '/api/auth/me', cookie=session(headers))[0], 401)

    def test_five_wrong_passwords_lock_the_email(self):
        signup('trava@x.com', role=self.role)
        for _ in range(5):
            self.assertEqual(call('POST', '/api/auth/login', {'email': 'trava@x.com', 'password': 'errada'})[1]['error'], 'E-mail ou senha incorretos.')
        status, body, _ = call('POST', '/api/auth/login', {'email': 'trava@x.com', 'password': 'senha-forte'})
        self.assertIn('Muitas tentativas', body['error'])

    def test_role_limits_tools(self):
        _, cookie = signup('ofertas@x.com', role=self.role)
        self.assertEqual(call('GET', '/api/jobs', cookie=cookie)[0], 403)
        self.assertEqual(call('POST', '/api/vector/trace', {}, cookie)[0], 403)
        self.assertEqual(call('GET', '/api/history', cookie=cookie)[0], 200)
        info = call('GET', '/api/info', cookie=cookie)[1]
        self.assertEqual([m['id'] for m in info['modules']], ['offers'])
        self.assertEqual(call('GET', '/api/admin/dados', cookie=cookie)[0], 403)

    def test_users_only_see_their_own_jobs_and_orders(self):
        role = call('POST', '/api/admin/cargo-salvar', {'name': 'Editor', 'modules': ['images', 'offers']}, self.admin_cookie)[1]['id']
        a, cookie_a = signup('a@x.com', role=role)
        _, cookie_b = signup('b@x.com', role=role)
        job = call('POST', '/api/jobs', {'kind': 'image'}, cookie_a)[1]
        self.assertEqual([j['id'] for j in call('GET', '/api/jobs', cookie=cookie_a)[1]], [job['id']])
        self.assertEqual(call('GET', '/api/jobs', cookie=cookie_b)[1], [])
        self.assertEqual(call('GET', '/api/job?id=' + job['id'], cookie=cookie_b)[1]['error'], 'Não encontrado.')
        self.assertEqual(call('POST', '/api/rename', {'id': job['id'], 'revision': 0, 'title': 'x'}, cookie_b)[1]['error'], 'Não encontrado.')
        self.assertEqual(call('GET', '/api/job?id=' + job['id'], cookie=cookie_a)[0], 200)
        self.assertEqual([e['action'] for e in call('GET', '/api/history', cookie=cookie_b)[1]], [])
        with connect() as db:
            db.execute("INSERT INTO ofertas_pedidos(id,template,titulo,dados,revision,criado,atualizado,owner) VALUES('p1','t','P','{}',0,?,?,?)", (now(), now(), a['id']))
        self.assertEqual(call('GET', '/api/ofertas/pedido?id=p1', cookie=cookie_b)[1]['error'], 'Não encontrado.')

    def test_superadmin_rules(self):
        status, body, _ = call('POST', '/api/admin/usuario-salvar', {'id': self.admin['id'], 'superadmin': False}, self.admin_cookie)
        self.assertEqual(status, 400)
        self.assertIn('último superadmin', body['error'])
        user, cookie = signup('sai@x.com', role=self.role)
        self.assertEqual(call('POST', '/api/admin/usuario-salvar', {'id': user['id'], 'superadmin': True}, cookie)[0], 403)
        self.assertEqual(call('POST', '/api/admin/usuario-salvar', {'id': user['id'], 'active': False}, self.admin_cookie)[0], 200)
        self.assertEqual(call('GET', '/api/auth/me', cookie=cookie)[0], 401)
        self.assertEqual(call('POST', '/api/admin/cargo-excluir', {'id': self.role}, self.admin_cookie)[0], 400)

    def test_password_reset_link(self):
        user, old_cookie = signup('esqueceu@x.com', role=self.role)
        self.assertEqual(call('POST', '/api/admin/redefinir-senha', {'id': user['id']}, old_cookie)[0], 403)
        status, body, _ = call('POST', '/api/admin/redefinir-senha', {'id': user['id']}, self.admin_cookie)
        token = body['link'].split('convite=')[1]
        self.assertEqual(call('GET', '/api/auth/convite?token=' + token)[1], {'email': 'esqueceu@x.com', 'reset': True})
        self.assertNotIn('esqueceu@x.com', [i['email'] for i in call('GET', '/api/admin/dados', cookie=self.admin_cookie)[1]['invites']])
        status, body, headers = call('POST', '/api/auth/cadastro', {'token': token, 'password': 'senha-nova-123'})
        self.assertEqual((status, body['id'], body['name']), (200, user['id'], 'Pessoa'))
        self.assertEqual(call('GET', '/api/auth/me', cookie=old_cookie)[0], 401)  # sessões antigas caem
        self.assertEqual(call('GET', '/api/auth/me', cookie=session(headers))[0], 200)
        self.assertEqual(call('POST', '/api/auth/login', {'email': 'esqueceu@x.com', 'password': 'senha-forte'})[0], 400)
        self.assertEqual(call('POST', '/api/auth/login', {'email': 'esqueceu@x.com', 'password': 'senha-nova-123'})[0], 200)
        self.assertEqual(call('POST', '/api/auth/cadastro', {'token': token, 'password': 'outra-senha-1'})[0], 400)  # uso único
        call('POST', '/api/admin/usuario-salvar', {'id': user['id'], 'active': False}, self.admin_cookie)
        self.assertIn('Reative', call('POST', '/api/admin/redefinir-senha', {'id': user['id']}, self.admin_cookie)[1]['error'])

    def test_my_account(self):
        user, cookie = signup('conta@x.com', 'Conta', role=self.role)
        other = session(call('POST', '/api/auth/login', {'email': 'conta@x.com', 'password': 'senha-forte'})[2])
        me = call('GET', '/api/auth/me', cookie=cookie)[1]
        self.assertEqual((me['roleName'], me['sessions']), ('Ofertas', 2))
        self.assertEqual(call('POST', '/api/auth/perfil', {'name': '  Conta Nova '}, cookie)[1]['name'], 'Conta Nova')
        self.assertEqual(call('POST', '/api/auth/perfil', {'name': ' '}, cookie)[0], 400)
        status, body, _ = call('POST', '/api/auth/senha', {'atual': 'errada', 'nova': 'outra-senha-1'}, cookie)
        self.assertEqual((status, body['error']), (400, 'A senha atual não confere.'))
        self.assertEqual(call('POST', '/api/auth/senha', {'atual': 'senha-forte', 'nova': 'outra-senha-1'}, cookie)[1], {'ok': True, 'removed': 1})
        self.assertEqual(call('GET', '/api/auth/me', cookie=other)[0], 401)   # outro aparelho caiu
        self.assertEqual(call('GET', '/api/auth/me', cookie=cookie)[0], 200)  # este continua
        self.assertEqual(call('POST', '/api/auth/login', {'email': 'conta@x.com', 'password': 'outra-senha-1'})[0], 200)
        self.assertEqual(call('POST', '/api/auth/sair-outros', {}, cookie)[1], {'ok': True, 'removed': 1})
        self.assertEqual(call('POST', '/api/auth/perfil', {'name': 'X'})[0], 401)

    def test_history_is_linked_to_user_and_tool(self):
        role = call('POST', '/api/admin/cargo-salvar', {'name': 'Relatório', 'modules': ['images', 'video', 'offers']}, self.admin_cookie)[1]['id']
        user, cookie = signup('relatorio@x.com', role=role)
        image = call('POST', '/api/jobs', {'kind': 'image'}, cookie)[1]
        video = call('POST', '/api/jobs', {'kind': 'video'}, cookie)[1]
        call('POST', '/api/rename', {'id': video['id'], 'revision': video['revision'], 'title': 'Vídeo do mês'}, cookie)
        with connect() as db: ofertas.registrar(db, 'Vídeo gerado', 'Teste', user['id'], render='r1')  # fila: fora da requisição
        with connect() as db:
            rows = [tuple(r) for r in db.execute('SELECT action, module FROM events WHERE owner=? ORDER BY id', (user['id'],))]
        self.assertEqual(rows, [('Trabalho criado', 'images'), ('Trabalho criado', 'video'), ('Edição renomeada', 'video'), ('Vídeo gerado', 'offers')])
        # "Apagar tudo" no Histórico só esconde da lista: o registro continua para os relatórios
        self.assertEqual(call('POST', '/api/history-delete', {'all': True}, cookie)[1]['count'], 4)
        self.assertEqual(call('GET', '/api/history', cookie=cookie)[1], [])
        with connect() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM events WHERE owner=? AND hidden=1', (user['id'],)).fetchone()[0], 4)

    def test_invite_link_from_admin(self):
        status, body, _ = call('POST', '/api/admin/convite', {'email': 'novo@x.com', 'role': self.role}, self.admin_cookie)
        self.assertEqual(status, 200)
        self.assertTrue(body['link'].startswith('http://localhost:8080/login.html?convite='))
        self.assertIn('novo@x.com', [i['email'] for i in call('GET', '/api/admin/dados', cookie=self.admin_cookie)[1]['invites']])


    def test_home_brand_tutorials_and_content_permissions(self):
        def raw(path, data, cookie, extra=None):
            path, _, qs = path.partition('?')
            env = {'REQUEST_METHOD': 'POST', 'PATH_INFO': path, 'QUERY_STRING': qs, 'HTTP_HOST': 'localhost:8080', 'HTTP_X_INDOOR': '1',
                   'CONTENT_LENGTH': str(len(data)), 'wsgi.input': io.BytesIO(data), 'HTTP_COOKIE': cookie}
            out = {}; body = b''.join(server.app(env, lambda s, h: out.update(status=s)))
            return int(out['status'][:3]), json.loads(body)
        roles = call('GET', '/api/admin/dados', cookie=self.admin_cookie)[1]['roles']
        eap_role = next(r for r in roles if r['brand'] == 'eap')
        self.assertEqual((eap_role['name'], sorted(eap_role['modules'])), ('Colaborador EAP', ['eap', 'ms6']))
        _, cookie = signup('eap@x.com', 'Marina Souza', eap_role['id'])
        info = call('GET', '/api/info', cookie=cookie)[1]
        self.assertEqual((info['brand'], sorted(m['id'] for m in info['modules'])), ('eap', ['eap', 'ms6']))
        self.assertEqual(call('GET', '/api/info', cookie=self.admin_cookie)[1]['brand'], 'indoor')
        image = io.BytesIO(); Image.new('RGB', (40, 30), 'red').save(image, 'PNG')
        self.assertEqual(raw('/api/home/foto', image.getvalue(), self.admin_cookie)[0], 200)
        mp4 = b'\x00\x00\x00\x18ftypmp42' + b'\x00' * 64
        for tool in ('offers', 'eap'): self.assertEqual(raw(f'/api/home/tutorial?tool={tool}', mp4, self.admin_cookie)[0], 200)
        self.assertEqual(raw('/api/home/tutorial?tool=eap', b'nao e video', self.admin_cookie)[0], 400)
        self.assertEqual(call('POST', '/api/admin/cargo-salvar', {'name': 'X', 'modules': [], 'brand': 'outra'}, self.admin_cookie)[0], 400)
        home = call('GET', '/api/home', cookie=cookie)[1]
        self.assertEqual(([t['tool'] for t in home['tutoriais']], len(home['fotos']), home['editar']), (['eap'], 1, False))
        self.assertEqual(raw('/api/home/foto', image.getvalue(), cookie)[0], 403)
        self.assertEqual(call('GET', '/home/tutorial/offers', cookie=cookie)[0], 403)
        self.assertEqual(call('GET', '/home/tutorial/eap', cookie=cookie)[0], 200)
        self.assertEqual(call('GET', home['fotos'][0]['url'], cookie=cookie)[0], 200)
        admin_home = call('GET', '/api/home', cookie=self.admin_cookie)[1]
        self.assertEqual(([t['tool'] for t in admin_home['tutoriais']], admin_home['editar']), (['offers', 'eap'], True))
        self.assertEqual(call('POST', '/api/home/foto-remover', {'id': home['fotos'][0]['id']}, self.admin_cookie)[1]['fotos'], [])


def clockify_entry(user, desc, tags, start, minutes, project='Cliente IC', running=False):
    end = None if running else (datetime.fromisoformat(start) + timedelta(minutes=minutes)).isoformat()
    return {'_id': uuid.uuid4().hex, 'description': desc, 'userName': user, 'projectName': project, 'clientName': '',
            'tags': [{'name': n} for n in tags], 'timeInterval': {'start': start, 'end': end, 'duration': None if running else minutes * 60}}


class RelatoriosTests(unittest.TestCase):
    """Relatório mensal a partir do Clockify (só superadmin) e o cronômetro próprio, guardado atrás de CRONOMETRO."""
    @classmethod
    def setUpClass(cls):
        server.AUTH = True
        server.CRONOMETRO = True
        cls.admin, cls.admin_cookie = signup('relatorios-chefe@x.com', 'Chefe Relatórios', superadmin=True)
        role = call('POST', '/api/admin/cargo-salvar', {'name': 'Edição', 'modules': ['images', 'video']}, cls.admin_cookie)[1]['id']
        cls.ana, cls.ana_cookie = signup('ana@x.com', 'Ana Lima', role)
        cls.bia, cls.bia_cookie = signup('bia@x.com', 'Bia Rocha', role)

    @classmethod
    def tearDownClass(cls): server.AUTH = False; server.CRONOMETRO = False

    def test_timer_stays_off_by_default(self):
        server.CRONOMETRO = False
        try:
            self.assertFalse(call('GET', '/api/info', cookie=self.ana_cookie)[1]['registraTempo'])
            self.assertEqual(call('GET', '/api/atividades/atual', cookie=self.ana_cookie)[1], {'atividade': None, 'registra': False})
        finally: server.CRONOMETRO = True

    def test_clockify_monthly_report_and_exports(self):
        entradas = [
            clockify_entry('Ana Lima', 'Veiculação - Clivet', ['Edição VT', 'Rende receita'], '2026-09-02T09:00:00-03:00', 90),
            clockify_entry('Ana Lima', 'Veiculação - Clivet', ['Edição VT'], '2026-09-02T14:00:00-03:00', 30),
            clockify_entry('Bia Rocha', 'Veiculação - Clivet (Alteração 1)', ['Edição VT'], '2026-09-03T10:00:00-03:00', 20),
            clockify_entry('Bia Rocha', 'veiculação - clivet (alteração 2)', [], '2026-09-04T10:00:00-03:00', 10),
            clockify_entry('Bia Rocha', 'Totem - Faculdade', ['Alteração de arte', 'Totem'], '2026-09-05T10:00:00-03:00', 15, 'Institucional'),
            clockify_entry('Ana Lima', '', [], '2026-09-30T23:00:00-03:00', 0, None, running=True),
        ]
        chamadas = []
        def falso(mes): chamadas.append(mes); return entradas
        with mock.patch.object(clockify, 'buscar', falso):
            clockify._cache.clear()
            self.assertEqual(call('GET', '/api/relatorios?mes=2026-09', cookie=self.ana_cookie)[0], 403)
            self.assertEqual(call('GET', '/relatorios/relatorio-2026-09.xlsx', cookie=self.ana_cookie)[0], 403)
            rel = call('GET', '/api/relatorios?mes=2026-09', cookie=self.admin_cookie)[1]
            t = rel['totais']
            self.assertEqual((t['registros'], t['tarefas'], t['alteracoes'], t['colaboradores'], t['emAndamento']), (6, 5, 3, 2, 1))
            self.assertAlmostEqual(t['segundos'], 225 * 60, delta=1)  # o cronômetro esquecido ligado conta só até o fim do mês
            self.assertEqual(rel['alteracoesPorTarefa'][0] | {'segundos': 0}, {'nome': 'Veiculação - Clivet', 'alteracoes': 2, 'segundos': 0})
            vt = next(x for x in rel['porEtiqueta'] if x['nome'] == 'Edição VT')
            self.assertEqual((vt['segundos'], vt['tarefas'], vt['alteracoes'], vt['cor']), (140 * 60, 2, 1, '#6A4FD1'))
            self.assertEqual([x['nome'] for x in rel['porProjeto']], ['Cliente IC', 'Sem projeto', 'Institucional'])
            self.assertEqual(chamadas, ['2026-09'])
            call('GET', '/api/relatorios?mes=2026-09', cookie=self.admin_cookie)
            self.assertEqual(chamadas, ['2026-09'])  # guardado por 2 minutos
            call('GET', '/api/relatorios?mes=2026-09&atualizar=1', cookie=self.admin_cookie)
            self.assertEqual(len(chamadas), 2)  # "Atualizar" busca de novo
            status, xlsx, headers = call('GET', '/relatorios/relatorio-2026-09.xlsx', cookie=self.admin_cookie)
            self.assertEqual((status, xlsx[:2]), (200, b'PK')); self.assertIn('attachment', headers['Content-Disposition'])
            planilha = load_workbook(io.BytesIO(xlsx))
            self.assertEqual(planilha['Registros'].max_row, 7)
            status, pdf, _ = call('GET', '/relatorios/relatorio-2026-09.pdf', cookie=self.admin_cookie)
            self.assertEqual((status, pdf[:5]), (200, b'%PDF-'))
            self.assertEqual(call('GET', '/api/relatorios?mes=2026-13', cookie=self.admin_cookie)[0], 400)
        self.assertEqual(clockify.intervalo('2026-09'), ('2026-09-01T03:00:00.000Z', '2026-10-01T02:59:59.999Z'))
        self.assertEqual(clockify.intervalo('2026-12')[1], '2027-01-01T02:59:59.999Z')
        clockify._cache.clear()
        with mock.patch.dict(os.environ, {'CLOCKIFY_API_KEY': ''}):
            status, erro, _ = call('GET', '/api/relatorios?mes=2026-08', cookie=self.admin_cookie)
        self.assertEqual(status, 400); self.assertIn('CLOCKIFY_API_KEY', json.dumps(erro, ensure_ascii=False))

    def test_timer_repetition_and_corrections(self):
        info = call('GET', '/api/info', cookie=self.ana_cookie)[1]
        self.assertTrue(info['registraTempo'])
        atual = call('GET', '/api/atividades/atual', cookie=self.ana_cookie)[1]
        self.assertIsNone(atual['atividade']); self.assertEqual(len(atual['etiquetas']), 10)
        self.assertEqual(call('POST', '/api/atividades/iniciar', {'cliente': 'Vantagem', 'tarefa': 'Ofertas', 'etiquetas': []}, self.ana_cookie)[0], 400)
        tarefa = {'cliente': 'Vantagem Supermercado', 'tarefa': 'Ofertas 11 a 14/10', 'etiquetas': ['adaptacao-encartes', 'edicao-vt', 'inexistente']}
        status, data, _ = call('POST', '/api/atividades/iniciar', tarefa, self.ana_cookie)
        self.assertEqual((status, data['atividade']['etiquetas'], data['atividade']['status']), (200, ['adaptacao-encartes', 'edicao-vt'], 'andamento'))
        original = data['atividade']['id']
        self.assertEqual(call('POST', '/api/atividades/iniciar', tarefa, self.ana_cookie)[0], 400)  # uma aberta por vez
        self.assertEqual(call('POST', '/api/atividades/pausar', {}, self.ana_cookie)[1]['atividade']['status'], 'pausada')
        self.assertEqual(call('POST', '/api/atividades/retomar', {}, self.ana_cookie)[1]['atividade']['status'], 'andamento')
        self.assertEqual(call('POST', '/api/atividades/corrigir', {'id': original, 'minutos': 30, 'motivo': 'x'}, self.ana_cookie)[0], 400)
        self.assertEqual(call('POST', '/api/atividades/finalizar', {}, self.ana_cookie)[1]['atividade']['status'], 'finalizada')
        # outra pessoa recebe o mesmo encarte de novo, com espaços e maiúsculas diferentes: é alteração
        repetida = call('POST', '/api/atividades/verificar', {'cliente': ' vantagem  supermercado', 'tarefa': 'OFERTAS 11 a 14/10'}, self.bia_cookie)[1]['repeticao']
        self.assertEqual((repetida['id'], repetida['pessoa']), (original, 'Ana Lima'))
        self.assertIsNone(call('POST', '/api/atividades/verificar', {'cliente': 'Vantagem Supermercado', 'tarefa': 'Ofertas 15 a 18/10'}, self.bia_cookie)[1]['repeticao'])
        segunda = call('POST', '/api/atividades/iniciar', tarefa | {'etiquetas': ['adaptacao-encartes'], 'alteracaoDe': original}, self.bia_cookie)[1]['atividade']
        self.assertEqual(segunda['alteracaoDe'], original)
        call('POST', '/api/atividades/finalizar', {}, self.bia_cookie)
        sugestoes = call('GET', '/api/atividades/sugestoes', cookie=self.bia_cookie)[1]
        self.assertIn('Vantagem Supermercado', sugestoes['clientes'])
        # correção: só a própria, com motivo; registrada
        self.assertEqual(call('POST', '/api/atividades/corrigir', {'id': original, 'minutos': 45, 'motivo': 'esqueci'}, self.bia_cookie)[0], 400)
        corrigida = call('POST', '/api/atividades/corrigir', {'id': original, 'minutos': 45, 'motivo': 'Esqueci de pausar no almoço'}, self.ana_cookie)[1]['atividade']
        self.assertEqual((corrigida['segundos'], corrigida['corrigido']), (2700, True))
        self.assertEqual(len(call('GET', '/api/atividades/minhas', cookie=self.ana_cookie)[1]['atividades']), 1)
        grafada = call('POST', '/api/atividades/iniciar', {'cliente': 'vantagem SUPERMERCADO', 'tarefa': 'Encarte novo', 'etiquetas': ['totem']}, self.ana_cookie)[1]['atividade']
        self.assertEqual(grafada['cliente'], 'Vantagem Supermercado')  # reaproveita a grafia já usada
        call('POST', '/api/atividades/finalizar', {}, self.ana_cookie)
        # o relatório interno do cronômetro guardado continua coerente
        rel = relatorios.relatorio(data['atividade']['inicio'][:7])
        self.assertGreaterEqual(rel['totais']['atividades'], 2); self.assertGreaterEqual(rel['totais']['alteracoes'], 1)
        encartes = next(x for x in rel['porEtiqueta'] if x['nome'] == 'Adaptação de encartes')
        self.assertGreaterEqual((encartes['atividades'], encartes['alteracoes']), (2, 1))

    def test_eap_role_does_not_track_time(self):
        roles = call('GET', '/api/admin/dados', cookie=self.admin_cookie)[1]['roles']
        _, cookie = signup('eap-tempo@x.com', 'Carlos EAP', next(r for r in roles if r['brand'] == 'eap')['id'])
        self.assertFalse(call('GET', '/api/info', cookie=cookie)[1]['registraTempo'])
        self.assertEqual(call('GET', '/api/atividades/atual', cookie=cookie)[1], {'atividade': None, 'registra': False})


class LocalModeTests(unittest.TestCase):
    def test_without_auth_everything_stays_open(self):
        server.AUTH = False
        self.assertEqual(call('GET', '/api/jobs')[0], 200)
        self.assertEqual(call('GET', '/')[0], 200)


if __name__ == '__main__':
    unittest.main()
